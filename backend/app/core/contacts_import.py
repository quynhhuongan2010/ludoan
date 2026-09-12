"""Doc file danh ba (.xlsx / .csv) -> (headers, rows) va doan cac truong chuan.

Khong phu thuoc pandas; dung openpyxl cho .xlsx va module csv chuan cho .csv.
"""

import csv
import io

from fastapi import HTTPException, status
from openpyxl import load_workbook

MAX_ROWS = 20000

_NAME_KEYS = ("họ và tên", "họ tên", "ho ten", "hoten", "full name", "fullname", "tên", "name")
_UNIT_KEYS = (
    "đơn vị công tác",
    "đơn vị",
    "don vi",
    "cơ quan",
    "co quan",
    "phòng ban",
    "phòng",
    "phong",
    "unit",
)
_POS_KEYS = (
    "cấp bậc chức vụ",
    "chức vụ",
    "chuc vu",
    "cấp bậc",
    "cap bac",
    "chức danh",
    "chuc danh",
    "position",
    "title",
)
_PHONE_KEYS = (
    "di động",
    "di dong",
    "số điện thoại",
    "so dien thoai",
    "điện thoại",
    "dien thoai",
    "sđt",
    "sdt",
    "số máy",
    "so may",
    "phone",
    "mobile",
    "đt",
)
_EMAIL_KEYS = ("email", "e-mail", "thư điện tử", "thu dien tu", "hòm thư", "hom thu", "mail")


def _norm(s) -> str:
    return (str(s) if s is not None else "").strip().lower()


def _match(header: str, keys) -> bool:
    h = _norm(header)
    return any(k in h for k in keys)


def map_standard_fields(row: dict) -> dict:
    """Doan cac truong chuan tu 1 dong {header: value}. Tra ve dict co the co gia tri None."""
    out = {"full_name": None, "unit": None, "position": None, "phone": None, "email": None}

    # phone: uu tien cot co chu 'di dong'
    phone_cols = [h for h in row if _match(h, _PHONE_KEYS)]
    phone_cols.sort(key=lambda x: 0 if ("di dong" in _norm(x) or "di động" in _norm(x)) else 1)
    for h in phone_cols:
        v = str(row[h]).strip()
        if v:
            out["phone"] = v
            break

    for h, v in row.items():
        val = str(v).strip() if v is not None else ""
        if not val:
            continue
        if out["full_name"] is None and _match(h, _NAME_KEYS):
            out["full_name"] = val
        elif out["unit"] is None and _match(h, _UNIT_KEYS):
            out["unit"] = val
        elif out["position"] is None and _match(h, _POS_KEYS):
            out["position"] = val
        elif out["email"] is None and _match(h, _EMAIL_KEYS):
            out["email"] = val
    return out


def _clean_headers(raw: list) -> list[str]:
    out: list[str] = []
    seen: dict[str, int] = {}
    for i, h in enumerate(raw):
        name = str(h).strip() if h is not None else ""
        if not name:
            name = f"Cột {i + 1}"
        if name in seen:
            seen[name] += 1
            name = f"{name} ({seen[name]})"
        else:
            seen[name] = 1
        out.append(name)
    return out


def _row_has_data(cells) -> bool:
    return any(c is not None and str(c).strip() for c in cells)


def _parse_xlsx(data: bytes):
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST, detail=f"Không đọc được file Excel: {e}"
        )
    ws = wb.active
    it = ws.iter_rows(values_only=True)

    header_row = None
    for r in it:
        if r is not None and _row_has_data(r):
            header_row = r
            break
    if header_row is None:
        wb.close()
        return [], []

    headers = _clean_headers(list(header_row))
    rows: list[dict] = []
    for r in it:
        if r is None or not _row_has_data(r):
            continue
        rec = {}
        for idx, h in enumerate(headers):
            v = r[idx] if idx < len(r) else None
            rec[h] = "" if v is None else str(v).strip()
        rows.append(rec)
    wb.close()
    return headers, rows


def _parse_csv(data: bytes):
    text = None
    for enc in ("utf-8-sig", "utf-8", "cp1258", "latin-1"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Không giải mã được file CSV (lưu lại dạng UTF-8 rồi thử lại)",
        )
    all_rows = list(csv.reader(io.StringIO(text)))
    if not all_rows:
        return [], []
    headers = _clean_headers(all_rows[0])
    rows: list[dict] = []
    for row in all_rows[1:]:
        if not any(str(c).strip() for c in row):
            continue
        rec = {h: (row[i].strip() if i < len(row) else "") for i, h in enumerate(headers)}
        rows.append(rec)
    return headers, rows


def parse_contacts_file(filename: str, data: bytes) -> tuple[list[str], list[dict]]:
    ext = filename.lower().rsplit(".", 1)[-1] if "." in (filename or "") else ""
    if ext in ("xlsx", "xlsm"):
        headers, rows = _parse_xlsx(data)
    elif ext == "csv":
        headers, rows = _parse_csv(data)
    else:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Chỉ hỗ trợ file danh bạ định dạng .xlsx hoặc .csv",
        )
    if not headers:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST, detail="File không có dòng tiêu đề cột"
        )
    if not rows:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST, detail="File không có dòng dữ liệu nào"
        )
    if len(rows) > MAX_ROWS:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"File có {len(rows)} dòng, vượt giới hạn {MAX_ROWS} dòng mỗi lần nhập",
        )
    return headers, rows
