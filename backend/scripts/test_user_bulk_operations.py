"""Kiem thu tu dong toan dien cho Tinh nang Quan ly Nguoi dung nang cao (v7.9.0):
1. Admin toan quyen xoa tai khoan chi huy (role 1, 2, 3)
2. Sinh file mau Excel (.xlsx) va Word (.docx)
3. Nhap nguoi dung hang loat tu file Excel (.xlsx)
4. Nhap nguoi dung hang loat tu file Word (.docx)
5. Xu ly loi khi nhap (trung username, sai dinh dang)
6. Don dep sach tai khoan thu nghiem (purge_test_users), giu lai admin
7. Khoi tao lai bo tai khoan mau Lu doan 21 day du
"""

import io
import pathlib
import sys

if sys.stdout.encoding != "utf-8":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent.parent))

import docx
import openpyxl
from app.core.database import SessionLocal
from app.models.unit import Unit
from app.models.user import User
from app.services import user_service
from scripts.seed_military_accounts import seed_military_accounts


def run_tests():
    print("=== BẮT ĐẦU KIỂM THỬ QUẢN LÝ NGƯỜI DÙNG & BULK OPERATIONS (v7.9.0) ===\n")

    with SessionLocal() as db:
        admin_user = db.query(User).filter(User.username == "admin").first()
        assert admin_user is not None, "Tài khoản admin phải tồn tại"

        # --- Test 1: Admin toàn quyền xoá tài khoản chỉ huy ---
        print("--- Test 1: Admin toàn quyền xoá tài khoản chỉ huy (Role 1/2/3) ---")
        # Tạo 1 tài khoản chỉ huy tạm thời để test xoá
        unit_bch = db.query(Unit).filter(Unit.name.like("%Ban chỉ huy%")).first()
        temp_commander = User(
            username="temp_commander_test",
            full_name="Chỉ huy Thử Nghiệm",
            rank="Đại tá",
            position="Phó Lữ đoàn trưởng",
            unit_id=unit_bch.id if unit_bch else None,
            role=2,  # Cấp phó lữ đoàn
            is_active=True,
            clearance=True,
            hashed_password="hash",
        )
        db.add(temp_commander)
        db.commit()
        db.refresh(temp_commander)

        # Admin tiến hành xoá
        user_service.delete_user(db, temp_commander.id, current_user=admin_user)
        check_deleted = db.query(User).filter(User.username == "temp_commander_test").first()
        assert check_deleted is None, "Tài khoản chỉ huy tạm phải bị xoá thành công bởi admin"
        print("-> Admin đã xoá thành công tài khoản cấp chỉ huy mà không bị chặn!")

        # --- Test 2: Sinh file mẫu Excel và Word ---
        print("\n--- Test 2: Sinh file mẫu Excel (.xlsx) và Word (.docx) ---")
        xlsx_bytes, xlsx_name, xlsx_mime = user_service.generate_user_template("excel")
        assert len(xlsx_bytes) > 1000, "File mẫu Excel phải có dung lượng hợp lệ"
        assert xlsx_name.endswith(".xlsx")
        assert "spreadsheet" in xlsx_mime
        print(f"-> Sinh file mẫu Excel thành công: {xlsx_name} ({len(xlsx_bytes)} bytes)")

        docx_bytes, docx_name, docx_mime = user_service.generate_user_template("word")
        assert len(docx_bytes) > 1000, "File mẫu Word phải có dung lượng hợp lệ"
        assert docx_name.endswith(".docx")
        assert "wordprocessingml" in docx_mime
        print(f"-> Sinh file mẫu Word thành công: {docx_name} ({len(docx_bytes)} bytes)")

        # --- Test 3: Nhập người dùng hàng loạt từ file Excel ---
        print("\n--- Test 3: Nhập người dùng hàng loạt từ file Excel (.xlsx) ---")
        # Dọn dẹp trước nếu các user test đã tồn tại từ lần chạy trước
        for test_un in ("test_soldier_1", "test_c5_leader", "test_tech_word"):
            old_u = db.query(User).filter(User.username == test_un).first()
            if old_u:
                user_service._clean_user_foreign_keys(db, old_u.id)
                db.delete(old_u)
        db.commit()

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["STT", "Họ và tên", "Tên đăng nhập", "Mật khẩu", "Cấp bậc", "Chức danh", "Đơn vị", "Vai trò"])
        ws.append([1, "Chiến sĩ Test 1", "test_soldier_1", "LuDoan21@2026", "Binh nhất", "Chiến sĩ TTLL", "Tiểu đoàn 1", "Cá nhân"])
        ws.append([2, "Đại đội trưởng Test 2", "test_c5_leader", "LuDoan21@2026", "Đại uý", "Đại đội trưởng", "Đại đội 5", "Chỉ huy đơn vị"])
        excel_buf = io.BytesIO()
        wb.save(excel_buf)

        res_excel = user_service.import_users_from_file(
            db, excel_buf.getvalue(), "danh_sach_test.xlsx", current_user=admin_user
        )
        assert res_excel["success_count"] == 2, f"Phải nhập thành công 2 user từ Excel, thực tế: {res_excel}"
        assert "test_soldier_1" in res_excel["created_usernames"]
        assert "test_c5_leader" in res_excel["created_usernames"]
        print(f"-> Nhập thành công {res_excel['success_count']} tài khoản từ file Excel!")

        # --- Test 4: Nhập người dùng hàng loạt từ file Word ---
        print("\n--- Test 4: Nhập người dùng hàng loạt từ file Word (.docx) ---")
        doc = docx.Document()
        tbl = doc.add_table(rows=1, cols=8)
        headers = ["STT", "Họ và tên", "Tên đăng nhập", "Mật khẩu", "Cấp bậc", "Chức danh", "Đơn vị", "Vai trò"]
        for i, h in enumerate(headers):
            tbl.rows[0].cells[i].text = h
        row1 = tbl.add_row().cells
        for i, val in enumerate([1, "Kỹ thuật viên Word", "test_tech_word", "LuDoan21@2026", "Thiếu uý QNCN", "Kỹ thuật viên TT2", "Trung tâm 2", "Cá nhân"]):
            row1[i].text = str(val)
        word_buf = io.BytesIO()
        doc.save(word_buf)

        res_word = user_service.import_users_from_file(
            db, word_buf.getvalue(), "danh_sach_test.docx", current_user=admin_user
        )
        assert res_word["success_count"] == 1, f"Phải nhập thành công 1 user từ Word, thực tế: {res_word}"
        assert "test_tech_word" in res_word["created_usernames"]
        print(f"-> Nhập thành công {res_word['success_count']} tài khoản từ file Word!")

        # --- Test 5: Bắt lỗi khi nhập trùng username ---
        print("\n--- Test 5: Kiểm tra bắt lỗi khi tên đăng nhập đã tồn tại ---")
        res_dup = user_service.import_users_from_file(
            db, excel_buf.getvalue(), "danh_sach_dup.xlsx", current_user=admin_user
        )
        assert res_dup["error_count"] == 2, "Cả 2 dòng phải bị báo lỗi trùng lặp"
        assert len(res_dup["errors"]) == 2
        print(f"-> Bắt lỗi trùng lặp chính xác: {res_dup['errors'][0]['error']}")

        # --- Test 6: Xoá sạch toàn bộ tài khoản thử nghiệm (Purge) ---
        print("\n--- Test 6: Xoá sạch toàn bộ tài khoản thử nghiệm (Purge test users) ---")
        total_before = db.query(User).count()
        purge_res = user_service.purge_test_users(db, current_user=admin_user)
        total_after = db.query(User).count()
        print(f"-> Đã dọn dẹp: {purge_res['purged_count']} tài khoản. Trước: {total_before}, Sau: {total_after}")
        assert total_after == 1, f"Sau purge chỉ còn lại đúng 1 tài khoản admin, thực tế: {total_after}"
        admin_check = db.query(User).first()
        assert admin_check.username == "admin", "Tài khoản còn lại phải là admin"
        print("-> Toàn bộ tài khoản thử nghiệm đã được dọn sạch, tài khoản Admin bảo toàn tuyệt đối!")

        # --- Test 7: Khởi tạo lại bộ 18 tài khoản chuẩn Lữ đoàn 21 ---
        print("\n--- Test 7: Tái lập bộ 18 tài khoản chuẩn biên chế Lữ đoàn 21 ---")
        seed_military_accounts()
        db.commit()  # Refresh snapshot de nhin thay thay doi tu session khac
        total_seeded = db.query(User).count()
        assert total_seeded >= 19, f"Tổng số tài khoản sau khi seed phải >= 19 (1 admin + 18 quân nhân), thực tế: {total_seeded}"
        print(f"-> Khởi tạo lại thành công đầy đủ {total_seeded} tài khoản chuẩn biên chế!")

    print("\n=== TOÀN BỘ KIỂM THỬ QUẢN LÝ NGƯỜI DÙNG & BULK OPERATIONS (v7.9.0) THÀNH CÔNG 100%! ===")


if __name__ == "__main__":
    run_tests()
