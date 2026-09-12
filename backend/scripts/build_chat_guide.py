# -*- coding: utf-8 -*-
"""Bien soan file Word: HUONG_DAN_CHAT_NOI_BO.docx

Tai lieu GIOI THIEU + HUONG DAN su dung KENH TIN NHAN TAC CHIEN NOI BO
(INTRA-CHAT) - phien ban hop dong API v8.1.0, day du toan bo tinh nang:
chat 1-1, chat nhom + luong duyet, tra loi/trich dan, sua tin, thu hoi, tha
cam xuc, ghim tin, chuyen tiep, tim kiem, tat thong bao, luu tru, bien nhan da
xem, "dang soan tin", truc tuyen/ngoai tuyen, doi ten nhom, phong QTV nhom,
tin he thong, quan tri thanh vien.

Hinh minh hoa = anh chup THAT tu he thong dang chay, sinh boi:
    venv/Scripts/python.exe scripts/capture_chat_guide.py

Chay (can PYTHONIOENCODING=utf-8):
    set PYTHONIOENCODING=utf-8
    venv/Scripts/python.exe scripts/build_chat_guide.py

Ket qua: docs/HUONG_DAN_CHAT_NOI_BO.docx
"""

import datetime
import pathlib

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = pathlib.Path(__file__).resolve().parent.parent.parent
SHOTS = ROOT / "docs" / "screenshots_khai_thac"
OUT = ROOT / "docs" / "HUONG_DAN_CHAT_NOI_BO.docx"

API_VERSION = "v8.1.0"

GREEN = RGBColor(0x22, 0x66, 0x37)
GRAYTXT = RGBColor(0x66, 0x66, 0x66)
REDTXT = RGBColor(0xC0, 0x39, 0x2B)


def _shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def _page_field(paragraph):
    run = paragraph.add_run()
    for kind, txt in [("begin", None), ("instrText", " PAGE "), ("separate", None), ("t", "1"), ("end", None)]:
        el = OxmlElement(f"w:{kind}")
        if kind == "instrText":
            el.set(qn("xml:space"), "preserve")
            el.text = txt
        elif kind == "t":
            el.text = txt
        else:
            el.set(qn("w:fldCharType"), kind)
        run._r.append(el)


class M:
    def __init__(self):
        d = Document()
        self.d = d
        st = d.styles["Normal"]
        st.font.name = "Times New Roman"
        st.font.size = Pt(12)
        for i, sz in [(1, 16), (2, 13.5), (3, 12.5)]:
            h = d.styles[f"Heading {i}"]
            h.font.name = "Times New Roman"
            h.font.size = Pt(sz)
            h.font.color.rgb = GREEN
            h.font.bold = True
        sec = d.sections[0]
        sec.page_width = Inches(8.27)
        sec.page_height = Inches(11.69)
        for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(sec, m, Inches(0.8))
        fp = sec.footer.paragraphs[0]
        fp.text = "Hướng dẫn Kênh Tin nhắn Tác chiến nội bộ – Lữ đoàn Thông tin 21   |   Trang "
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in fp.runs:
            r.font.size = Pt(9)
            r.font.color.rgb = GRAYTXT
        _page_field(fp)

    def h(self, level, text):
        self.d.add_heading(text, level=level)

    def p(self, text, bold=False, italic=False, color=None, size=None):
        par = self.d.add_paragraph()
        r = par.add_run(text)
        r.bold = bold
        r.italic = italic
        if color:
            r.font.color.rgb = color
        if size:
            r.font.size = Pt(size)
        return par

    def steps(self, items):
        for it in items:
            self.d.add_paragraph(str(it), style="List Number")

    def bullets(self, items):
        for it in items:
            self.d.add_paragraph(str(it), style="List Bullet")

    def note(self, text, fill="FFF3CD", label="LƯU Ý"):
        t = self.d.add_table(rows=1, cols=1)
        c = t.cell(0, 0)
        _shade(c, fill)
        c.paragraphs[0].text = ""
        run = c.paragraphs[0].add_run(f"● {label}: ")
        run.bold = True
        c.paragraphs[0].add_run(text)
        for r in c.paragraphs[0].runs:
            r.font.size = Pt(11)
        self.d.add_paragraph()

    def figure(self, stem, caption):
        path = next(
            (SHOTS / f"{stem}{ext}" for ext in (".jpg", ".jpeg", ".png") if (SHOTS / f"{stem}{ext}").exists()),
            None,
        )
        if path is None:
            self.p(f"[Thiếu hình: {stem} — chạy scripts/capture_chat_guide.py]", italic=True, color=REDTXT)
            return
        max_w, max_h = Inches(6.6), Inches(8.4)
        w, h = max_w, None
        try:
            from PIL import Image

            with Image.open(path) as im:
                pw, ph = im.size
            h = int(max_w * ph / pw)
            if h > max_h:
                h = max_h
                w = int(max_h * pw / ph)
        except Exception:  # noqa: BLE001
            h = None
        if h:
            self.d.add_picture(str(path), width=w, height=h)
        else:
            self.d.add_picture(str(path), width=w)
        self.d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap = self.d.add_paragraph(f"Ảnh màn hình: {caption}")
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in cap.runs:
            r.italic = True
            r.font.size = Pt(10)
            r.font.color.rgb = GRAYTXT
        self.d.add_paragraph()

    def table(self, headers, rows, widths=None):
        t = self.d.add_table(rows=1, cols=len(headers))
        t.style = "Table Grid"
        for i, htext in enumerate(headers):
            cell = t.rows[0].cells[i]
            cell.text = ""
            run = cell.paragraphs[0].add_run(htext)
            run.bold = True
            run.font.size = Pt(10.5)
            _shade(cell, "E7ECEA")
        for row in rows:
            cells = t.add_row().cells
            for i, val in enumerate(row):
                cells[i].text = ""
                run = cells[i].paragraphs[0].add_run(str(val))
                run.font.size = Pt(10.5)
        if widths:
            for i, wv in enumerate(widths):
                for r in t.rows:
                    r.cells[i].width = Inches(wv)
        self.d.add_paragraph()

    def pagebreak(self):
        self.d.add_page_break()

    def save(self):
        self.d.save(str(OUT))


def build():
    m = M()
    d = m.d
    today = datetime.date.today().strftime("%d/%m/%Y")

    # ---- COVER ----
    for _ in range(3):
        d.add_paragraph()
    t = d.add_paragraph("HƯỚNG DẪN SỬ DỤNG")
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t.runs[0].bold = True
    t.runs[0].font.size = Pt(26)
    t.runs[0].font.color.rgb = GREEN
    s = d.add_paragraph("KÊNH TIN NHẮN TÁC CHIẾN NỘI BỘ (INTRA-CHAT)\nCỔNG THÔNG TIN ĐIỆN TỬ – LỮ ĐOÀN THÔNG TIN 21")
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    s.runs[0].font.size = Pt(14)
    s.runs[0].bold = True
    for _ in range(2):
        d.add_paragraph()
    box = d.add_table(rows=1, cols=1)
    box.style = "Table Grid"
    c = box.cell(0, 0)
    _shade(c, "F2F6F2")
    for line in [
        "Tài liệu giới thiệu & hướng dẫn khai thác cho Chỉ huy và cán bộ, quân nhân đơn vị",
        f"Phiên bản hợp đồng API: {API_VERSION} — có đầy đủ tính năng nâng cấp",
        f"Ngày biên soạn: {today}",
        "Hình minh hoạ là ẢNH CHỤP THẬT từ hệ thống đang chạy (không phải hình dựng)",
        "LƯU HÀNH NỘI BỘ",
    ]:
        pr = c.add_paragraph(line)
        pr.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pr.runs[0].font.size = Pt(12)
        if line == "LƯU HÀNH NỘI BỘ":
            pr.runs[0].bold = True
            pr.runs[0].font.color.rgb = REDTXT
    c.paragraphs[0].text = ""
    m.pagebreak()

    # ---- MUC LUC NHANH ----
    m.h(1, "MỤC LỤC NHANH")
    m.table(
        ["Phần", "Nội dung"],
        [
            ["1", "Vì sao đơn vị nên dùng kênh tin nhắn nội bộ này"],
            ["2", "Tổng quan giao diện 3 cột"],
            ["3", "Nhắn tin trực tiếp 1-1"],
            ["4", "Lập nhóm kíp trực & luồng duyệt nhóm"],
            ["5", "Soạn tin: mức độ KHẨN, biểu tượng, gửi ảnh / tài liệu / tệp nén"],
            ["6", "Trả lời – trích dẫn tin nhắn"],
            ["7", "Thanh công cụ trên từng tin nhắn (sửa · thu hồi · cảm xúc · ghim · chuyển tiếp)"],
            ["8", "Sửa tin nhắn & Thu hồi tin nhắn"],
            ["9", "Thả biểu tượng cảm xúc (reaction)"],
            ["10", "Ghim tin nhắn quan trọng"],
            ["11", "Chuyển tiếp tin nhắn sang hội thoại khác"],
            ["12", "Tìm kiếm trong hội thoại & tìm toàn bộ"],
            ["13", "Nhận tin thời gian thực: “đang soạn tin”, trực tuyến/ngoại tuyến, biên nhận “Đã xem”"],
            ["14", "Tin hệ thống (đổi tên nhóm, thêm/đưa thành viên ra khỏi nhóm)"],
            ["15", "Quản trị nhóm: đổi tên, phong/gỡ QTV nhóm, xoá thành viên, xoá nhóm"],
            ["16", "Tắt thông báo & Lưu trữ hội thoại"],
            ["17", "Phân quyền tóm tắt"],
            ["18", "An toàn thông tin & Nhật ký"],
            ["19", "Kết luận & đề xuất"],
        ],
        widths=[0.7, 5.9],
    )
    m.pagebreak()

    # =============================================================== 1
    m.h(1, "1. VÌ SAO ĐƠN VỊ NÊN DÙNG KÊNH TIN NHẮN NỘI BỘ NÀY")

    m.h(2, "1.1. Thực trạng đang gặp")
    m.p("Hiện nay việc trao đổi công việc, chỉ đạo – báo cáo, hiệp đồng kíp trực giữa Sở chỉ huy và "
        "các đầu mối phần lớn vẫn qua điện thoại cá nhân và các ứng dụng nhắn tin trên mạng công cộng "
        "(Zalo, Messenger, Viber…). Cách làm này bộc lộ nhiều hạn chế:")
    m.bullets([
        "Nội dung công việc, tài liệu, hình ảnh của đơn vị đi qua máy chủ của doanh nghiệp ngoài, "
        "phần lớn đặt ở nước ngoài — nguy cơ lộ, lọt thông tin, không kiểm soát được.",
        "Không phân quyền: ai cũng lập được nhóm, thêm/xoá người tuỳ ý; khó xác định trách nhiệm.",
        "Không lưu vết tập trung: khi cần rà soát, đối chiếu chỉ đạo thì phụ thuộc máy cá nhân từng người.",
        "Phụ thuộc Internet và nhà cung cấp: mất mạng ngoài, bị chặn dịch vụ, đổi số điện thoại… là gián đoạn.",
        "Lẫn lộn việc công – việc tư trên cùng một tài khoản, một ứng dụng.",
    ])

    m.h(2, "1.2. Kênh Tin nhắn Tác chiến nội bộ giải quyết như thế nào")
    m.table(
        ["Tiêu chí", "Nhắn tin qua mạng công cộng", "Kênh Tin nhắn Tác chiến nội bộ (phần mềm này)"],
        [
            ["Đường truyền", "Qua Internet, máy chủ doanh nghiệp ngoài", "CHỈ chạy trong mạng LAN nội bộ đơn vị, không cần Internet"],
            ["Nơi lưu dữ liệu", "Máy chủ nhà cung cấp (thường ở nước ngoài)", "Cơ sở dữ liệu đặt tại máy chủ của đơn vị"],
            ["Tài khoản", "SĐT/email cá nhân, lẫn việc công – tư", "Tài khoản do đơn vị cấp, gắn cấp bậc – chức danh – đơn vị"],
            ["Phân quyền", "Không", "6 mức vai trò; lập nhóm và duyệt nhóm theo cấp"],
            ["Lập nhóm", "Ai cũng lập, không kiểm soát", "Ban chỉ huy lập trực tiếp; cán bộ cấp dưới phải được duyệt"],
            ["Sửa / thu hồi tin", "Gần như không, hoặc chỉ trong ít phút", "Người gửi sửa/thu hồi bất kỳ lúc nào; chỉ huy & QTV nhóm thu hồi được tin sai của người khác"],
            ["Lưu vết", "Rải rác trên máy cá nhân", "Nhật ký an ninh tập trung: ai tạo/duyệt/từ chối/đổi tên/xoá nhóm, khi nào"],
            ["Tệp đính kèm", "Không kiểm tra", "Quét chữ ký nhị phân, chặn mã độc / web shell / tệp thực thi"],
            ["Chi phí / lệ thuộc", "Miễn phí nhưng lệ thuộc bên thứ ba", "Không lệ thuộc; đơn vị tự chủ hoàn toàn"],
        ],
        widths=[1.3, 2.3, 3.0],
    )

    m.h(2, "1.3. Điểm mạnh nổi bật (bản nâng cấp)")
    m.bullets([
        "Hoạt động 100% trong mạng nội bộ (LAN) — mất Internet vẫn dùng bình thường.",
        "Giao diện quen thuộc như Zalo / Messenger: danh sách hội thoại, khung chat, gửi ảnh/tệp, biểu tượng cảm xúc, đánh dấu KHẨN.",
        "ĐẦY ĐỦ thao tác trên từng tin nhắn: trả lời – trích dẫn, sửa, thu hồi, thả cảm xúc, ghim, chuyển tiếp.",
        "Nhận tin THỜI GIAN THỰC: tin mới, chỉnh sửa, thu hồi, cảm xúc, ghim đều hiện ngay; có dòng “đang soạn tin…”, chấm “đang trực tuyến”, biên nhận “Đã xem”.",
        "Quản trị nhóm rõ ràng: đổi tên nhóm, phong/gỡ Quản trị viên nhóm, mọi thay đổi thành viên đều để lại “tin hệ thống”.",
        "Tìm kiếm nhanh trong từng hội thoại và trên toàn bộ hội thoại của mình.",
        "Tắt thông báo và Lưu trữ từng hội thoại để danh sách gọn, tập trung việc chính.",
        "Phân quyền và luồng DUYỆT NHÓM giữ kiểm soát; mọi thao tác quan trọng đều ghi Nhật ký an ninh.",
    ])

    m.h(2, "1.4. Đã kiểm chứng bằng thao tác thật")
    m.p("Toàn bộ hình trong tài liệu này được chụp trực tiếp từ hệ thống đang chạy thật (Backend + "
        "cơ sở dữ liệu thật), điều khiển bằng kịch bản tự động hoá trình duyệt. Các nhóm tính năng đã "
        "được kiểm thử tự động 44 tình huống (sửa/thu hồi + chặn sai quyền, cảm xúc, ghim, chuyển tiếp, "
        "tìm kiếm, đổi tên nhóm sinh tin hệ thống, phong/gỡ QTV, tắt thông báo, lưu trữ, biên nhận…) — "
        "tất cả đạt yêu cầu.")
    m.figure("chat_01_bang_tin", "Đăng nhập bằng tài khoản do đơn vị cấp; nút “Tin nhắn” trên thanh trên cùng")
    m.pagebreak()

    # =============================================================== 2
    m.h(1, "2. TỔNG QUAN GIAO DIỆN")
    m.p("Vào mục “Tin nhắn” trên thanh trên cùng. Màn hình chia 3 cột:")
    m.bullets([
        "CỘT TRÁI — danh sách hội thoại: ô tìm kiếm; 5 thẻ lọc (Tất cả / Trực tiếp / Nhóm / Chưa đọc / "
        "Lưu trữ); nếu bạn có quyền duyệt, phía trên còn khối “Nhóm chờ duyệt”.",
        "CỘT GIỮA — khung chat: tên hội thoại, dòng trạng thái (số đồng chí trực tuyến), THANH TIN ĐÃ GHIM "
        "(nếu có), các bong bóng tin nhắn, và thanh công cụ soạn tin ở dưới.",
        "CỘT PHẢI — thông tin hội thoại: hồ sơ nhóm/đồng chí, hai nút nhanh “Tắt thông báo” / “Lưu trữ”, "
        "danh sách thành viên (kèm dấu QTV, chấm trực tuyến), các tệp đã gửi; với người tạo nhóm / Quản trị "
        "còn có nút đổi tên nhóm, xoá thành viên, xoá nhóm.",
        "Ba nút công cụ ở góc phải khung chat: KÍNH LÚP (tìm trong hội thoại) · CON MẮT GẠCH (tắt/bật "
        "thông báo) · CHỮ i (mở/đóng cột thông tin).",
    ])
    m.figure("chat_10_tong_quan", "Giao diện Kênh Tin nhắn Tác chiến nội bộ — bố cục 3 cột, có thanh tin đã ghim")
    m.pagebreak()

    # =============================================================== 3
    m.h(1, "3. NHẮN TIN TRỰC TIẾP 1-1")
    m.steps([
        "Bấm nút “+ Nhắn riêng” ở góc trên bên phải.",
        "Tìm và chọn đồng chí cần trao đổi (tìm theo họ tên, chức danh, đơn vị).",
        "Cửa sổ trò chuyện 1-1 mở ra; nhập nội dung, nhấn Enter để gửi.",
    ])
    m.figure("chat_18_nhan_rieng", "Chọn quân nhân để mở cuộc trò chuyện 1-1")
    m.figure("chat_20_chat_1_1", "Khung trò chuyện trực tiếp 1-1 (có cả tin được chuyển tiếp từ nhóm sang)")
    m.note("Cuộc trò chuyện 1-1 KHÔNG cần duyệt. Nội dung chỉ hiển thị cho đúng 2 người trong cuộc.")
    m.pagebreak()

    # =============================================================== 4
    m.h(1, "4. LẬP NHÓM KÍP TRỰC & LUỒNG DUYỆT NHÓM")

    m.h(2, "4.1. Ai được lập nhóm, ai duyệt")
    m.table(
        ["Vai trò tài khoản", "Lập nhóm", "Duyệt nhóm chờ"],
        [
            ["Quản trị hệ thống (0), Lữ trưởng – Chính uỷ (1)", "Trực tiếp — nhóm dùng được ngay", "CÓ"],
            ["Lữ phó – Phó chính uỷ (2), Chỉ huy đơn vị đầu mối (3)", "Trực tiếp — nhóm dùng được ngay", "Không"],
            ["Cán bộ / nhân viên (4), Người dùng (5)", "Gửi yêu cầu — nhóm ở trạng thái “Chờ duyệt”", "Không"],
        ],
        widths=[3.2, 2.0, 1.4],
    )

    m.h(2, "4.2. Lập nhóm")
    m.steps([
        "Bấm “+ Lập nhóm kíp trực”.",
        "Đặt TÊN NHÓM (ví dụ: “Kíp trực SSCĐ – Sở chỉ huy Lữ đoàn”).",
        "Tích chọn các đồng chí tham gia.",
        "Bấm “Tạo nhóm ngay”.",
    ])
    m.figure("chat_13_lap_nhom", "Biểu mẫu lập nhóm kíp trực")
    m.p("Nếu bạn thuộc Ban chỉ huy (vai trò 0–3): nhóm được tạo và dùng được ngay. Nếu bạn là cán bộ "
        "cấp dưới (vai trò 4–5): hệ thống báo “Đã gửi yêu cầu tạo nhóm — chờ chỉ huy Lữ đoàn duyệt”.")

    m.h(2, "4.3. Nhóm đang chờ duyệt — phía người tạo")
    m.p("Nhóm chờ duyệt CHỈ hiển thị với người tạo (các thành viên khác chưa nhìn thấy). Trong lúc "
        "chờ: chưa gửi được tin nhắn, chưa thêm được thành viên; ô nhập bị khoá kèm dòng thông báo.")
    m.figure("chat_19_nguoi_dung_cho_duyet", "Người tạo nhìn thấy nhóm của mình với nhãn “Chờ duyệt” và ô nhập bị khoá")

    m.h(2, "4.4. Chỉ huy duyệt hoặc từ chối")
    m.p("Tài khoản Quản trị / Lữ trưởng / Chính uỷ thấy khối “Nhóm chờ duyệt (N)” ở đầu cột trái, "
        "kèm nút “Duyệt” và “Từ chối” cho từng nhóm.")
    m.figure("chat_14_nhom_cho_duyet", "Khối “Nhóm chờ duyệt” với nút Duyệt / Từ chối")
    m.steps([
        "Bấm “Duyệt” → nhóm chuyển sang trạng thái dùng được; mọi thành viên bắt đầu trao đổi.",
        "Hoặc bấm “Từ chối” → nhập LÝ DO (bắt buộc) rồi xác nhận. Người tạo sẽ thấy nhóm bị từ chối kèm lý do và có thể tự xoá nhóm đó.",
    ])
    m.figure("chat_15_tu_choi_ly_do", "Từ chối nhóm — bắt buộc nhập lý do để người tạo nắm được")
    m.note("Mọi lần tạo / duyệt / từ chối / đổi tên nhóm đều được ghi vào Nhật ký an ninh (ai, nhóm nào, khi nào, lý do).")
    m.pagebreak()

    # =============================================================== 5
    m.h(1, "5. SOẠN TIN: MỨC ĐỘ KHẨN, BIỂU TƯỢNG & GỬI TỆP")
    m.p("Thanh công cụ phía dưới khung chat gồm: Tệp · Ảnh · Biểu cảm · nút Mức độ (Thường / KHẨN).")
    m.figure("chat_12_cong_cu_emoji", "Thanh công cụ soạn tin và bảng biểu tượng quân sự")
    m.h(2, "5.1. Gửi tin & đánh dấu KHẨN")
    m.steps([
        "Nhập nội dung vào ô soạn, nhấn Enter để gửi (Shift+Enter để xuống dòng).",
        "Bấm nút “Mức độ: Thường” để chuyển thành “MỨC ĐỘ: KHẨN” — tin gửi đi sẽ có nhãn “HỎA TỐC / KHẨN” nổi bật.",
        "Bấm “Biểu cảm” để chèn biểu tượng quân sự (🫡 🚩 🎖️ ⚡ 🛡️ …).",
    ])
    m.h(2, "5.2. Gửi ảnh, tài liệu, tệp nén")
    m.steps([
        "Bấm “Ảnh” để gửi hình (jpg/png/webp/gif) — hiển thị xem trước ngay trong khung chat, bấm để phóng to.",
        "Bấm “Tệp” để gửi tài liệu (.pdf .doc .docx .xls .xlsx .ppt .pptx), video (.mp4 .webm .mov…) hoặc TỆP NÉN (.zip .rar .7z .tar .gz .bz2).",
        "Người nhận thấy thẻ tệp kèm tên gốc, bấm để tải về. Cột phải liệt kê toàn bộ “Tệp tin đã gửi” của hội thoại.",
    ])
    m.figure("chat_11_khung_chat_tep", "Tin nhắn kèm ảnh và tệp nén; cột phải liệt kê “Tệp tin đã gửi”")
    m.note("Tệp tải lên được quét chữ ký nhị phân: từ chối tệp thực thi (.exe/PE, ELF), web shell, tệp "
           "nén/ảnh/tài liệu sai định dạng. Hạn mức: ảnh & tài liệu theo MAX_UPLOAD_MB; video & tệp nén "
           "theo MAX_VIDEO_UPLOAD_MB (đặt trong cấu hình máy chủ).")
    m.pagebreak()

    # =============================================================== 6
    m.h(1, "6. TRẢ LỜI – TRÍCH DẪN TIN NHẮN")
    m.p("Khi trao đổi nhiều nội dung đan xen, hãy trả lời đích danh một tin để người đọc biết bạn "
        "đang nói về việc nào.")
    m.steps([
        "Đưa chuột lên tin cần trả lời → thanh công cụ hiện ra → bấm biểu tượng “Trả lời” (mũi tên vòng).",
        "Phía trên ô soạn xuất hiện khối “Đang trả lời <tên đồng chí>” kèm trích đoạn nội dung.",
        "Nhập câu trả lời rồi gửi. Tin gửi đi mang theo khối trích dẫn; bấm vào khối đó sẽ nhảy tới tin gốc.",
        "Muốn bỏ trả lời: bấm dấu “X” trên khối “Đang trả lời”.",
    ])
    m.figure("chat_21_tra_loi_trich_dan", "Tin trả lời hiển thị kèm khối trích dẫn tin gốc")
    m.pagebreak()

    # =============================================================== 7
    m.h(1, "7. THANH CÔNG CỤ TRÊN TỪNG TIN NHẮN")
    m.p("Đưa chuột lên bất kỳ tin nhắn nào (chưa bị thu hồi), một thanh công cụ nhỏ hiện ra ngay cạnh "
        "bong bóng tin, gồm các nút:")
    m.table(
        ["Nút", "Tác dụng", "Ai dùng được"],
        [
            ["Trả lời", "Trả lời – trích dẫn đúng tin đó (xem phần 6)", "Mọi thành viên"],
            ["Thả cảm xúc", "Mở bảng cảm xúc nhanh để thả lên tin (xem phần 9)", "Mọi thành viên"],
            ["Chuyển tiếp", "Gửi tin đó sang hội thoại khác (xem phần 11)", "Mọi thành viên"],
            ["Ghim", "Ghim / bỏ ghim tin lên đầu khung chat (xem phần 10)", "QTV nhóm / Ban chỉ huy; chat 1-1: cả hai"],
            ["Sửa", "Sửa nội dung tin (xem phần 8)", "CHỈ người gửi tin đó"],
            ["Thu hồi", "Gỡ nội dung tin với mọi người (xem phần 8)", "Người gửi, hoặc QTV nhóm, hoặc Ban chỉ huy"],
        ],
        widths=[1.1, 3.7, 1.8],
    )
    m.figure("chat_27_thanh_cong_cu_tin_nhan", "Thanh công cụ hiện khi đưa chuột lên một tin nhắn")
    m.pagebreak()

    # =============================================================== 8
    m.h(1, "8. SỬA TIN NHẮN & THU HỒI TIN NHẮN")

    m.h(2, "8.1. Sửa tin")
    m.steps([
        "Đưa chuột lên tin CỦA MÌNH → bấm biểu tượng “Sửa” (bút chì).",
        "Nội dung tin chuyển thành ô nhập ngay tại chỗ; sửa xong bấm “Lưu” (hoặc Enter). Bấm “Huỷ” (hoặc Esc) để bỏ.",
        "Tin đã sửa hiển thị thêm chữ “(đã sửa)” ở dòng giờ để minh bạch.",
    ])
    m.p("Chỉ NGƯỜI GỬI mới sửa được tin của mình. Không sửa được tin đã thu hồi và tin hệ thống.", italic=True, color=GRAYTXT)

    m.h(2, "8.2. Thu hồi tin")
    m.steps([
        "Đưa chuột lên tin cần thu hồi → bấm biểu tượng “Thu hồi” (thùng rác) → xác nhận.",
        "Nội dung tin (kèm tệp đính kèm nếu có) bị gỡ với MỌI thành viên; chỗ đó chỉ còn dòng chữ nghiêng “Tin nhắn đã được thu hồi”.",
    ])
    m.p("Quyền thu hồi: NGƯỜI GỬI tin đó; HOẶC Quản trị viên nhóm; HOẶC Ban chỉ huy (vai trò 0–3). "
        "Người thường KHÔNG thu hồi được tin của người khác.")
    m.figure("chat_22_sua_thu_hoi", "Tin đã sửa mang nhãn “(đã sửa)”; tin đã thu hồi chỉ còn dòng “Tin nhắn đã được thu hồi”")
    m.note("Thu hồi là gỡ nội dung hiển thị và xoá tệp đính kèm của tin đó; dòng “đã thu hồi” vẫn giữ "
           "lại để không mất mạch hội thoại.")
    m.pagebreak()

    # =============================================================== 9
    m.h(1, "9. THẢ BIỂU TƯỢNG CẢM XÚC (REACTION)")
    m.p("Cảm xúc giúp xác nhận nhanh (“đã đọc”, “đồng ý”, “rõ”) mà không cần gửi thêm tin.")
    m.steps([
        "Đưa chuột lên tin → bấm biểu tượng “Thả cảm xúc” (mặt cười) → chọn một biểu tượng.",
        "Biểu tượng xuất hiện thành “chip” nhỏ ngay dưới tin, kèm số lượng người đã thả.",
        "Bấm lại vào chip đang sáng (của mình) để GỠ cảm xúc. Nhiều người có thể thả cùng lúc, nhiều loại khác nhau.",
    ])
    m.figure("chat_23_cam_xuc", "Các “chip” cảm xúc hiển thị dưới tin nhắn (👍 1 · 🫡 1 · ✅ 1)")
    m.pagebreak()

    # =============================================================== 10
    m.h(1, "10. GHIM TIN NHẮN QUAN TRỌNG")
    m.p("Ghim để đưa tin quan trọng (mệnh lệnh, kế hoạch, số điện thoại trực…) lên THANH GHIM ngay "
        "dưới tên hội thoại, ai vào cũng thấy đầu tiên.")
    m.steps([
        "Đưa chuột lên tin → bấm biểu tượng “Ghim”.",
        "Tin xuất hiện trên thanh ghim màu vàng nhạt ở đầu khung chat. Bấm vào dòng trên thanh ghim để nhảy tới tin gốc.",
        "Nhiều tin ghim: thanh hiện tin mới nhất và nút “+N” để mở rộng / “Thu gọn”.",
        "Bỏ ghim: đưa chuột lên tin → bấm lại biểu tượng “Ghim”.",
    ])
    m.p("Trong NHÓM: chỉ Quản trị viên nhóm hoặc Ban chỉ huy được ghim/bỏ ghim. Trong chat 1-1: cả hai bên đều ghim được.",
        italic=True, color=GRAYTXT)
    m.figure("chat_24_ghim_tin", "Thanh “tin đã ghim” hiển thị ngay dưới tên hội thoại")
    m.pagebreak()

    # =============================================================== 11
    m.h(1, "11. CHUYỂN TIẾP TIN NHẮN SANG HỘI THOẠI KHÁC")
    m.steps([
        "Đưa chuột lên tin cần chuyển → bấm biểu tượng “Chuyển tiếp” (máy bay giấy).",
        "Hộp thoại hiện danh sách các hội thoại KHÁC mà bạn đang tham gia — chọn hội thoại đích.",
        "Tin được đăng sang hội thoại đích, có nhãn “Chuyển tiếp từ <người gửi gốc>”.",
    ])
    m.p("Bạn phải là thành viên của CẢ hội thoại nguồn và hội thoại đích. Không chuyển tiếp được tin "
        "đã thu hồi hoặc tin hệ thống.", italic=True, color=GRAYTXT)
    m.figure("chat_25_chuyen_tiep", "Hộp thoại chọn hội thoại đích để chuyển tiếp tin nhắn")
    m.pagebreak()

    # =============================================================== 12
    m.h(1, "12. TÌM KIẾM")
    m.h(2, "12.1. Tìm trong một hội thoại")
    m.steps([
        "Mở hội thoại → bấm nút KÍNH LÚP ở góc phải khung chat.",
        "Gõ từ khoá → bấm “Tìm”. Kết quả liệt kê bên dưới: người gửi · nội dung · thời gian.",
        "Bấm một kết quả để nhảy tới đúng tin đó trong khung chat (tin được làm sáng vài giây).",
    ])
    m.figure("chat_26_tim_trong_hoi_thoai", "Bảng tìm kiếm trong hội thoại và danh sách kết quả")
    m.h(2, "12.2. Tìm trên toàn bộ hội thoại")
    m.p("Ngoài ra hệ thống còn hỗ trợ tìm nhanh trên toàn bộ hội thoại mà bạn tham gia (bỏ qua các "
        "tin đã thu hồi) — dùng khi không nhớ nội dung nằm ở nhóm nào.")
    m.pagebreak()

    # =============================================================== 13
    m.h(1, "13. NHẬN TIN THỜI GIAN THỰC")
    m.bullets([
        "Tin nhắn mới, tin được SỬA, tin bị THU HỒI, CẢM XÚC vừa thả, tin vừa GHIM — tất cả hiện ngay cho mọi thành viên đang mở kênh, KHÔNG cần bấm tải lại (F5).",
        "Dòng “<tên đồng chí> đang soạn tin…” hiện ngay trên ô soạn khi có người trong hội thoại đang gõ.",
        "Chấm xanh “đang trực tuyến” hiển thị ở danh sách hội thoại, ở đầu khung chat và cạnh từng thành viên trong cột phải; khi đối phương thoát ra thì chuyển “Ngoại tuyến”.",
        "Biên nhận “✓✓ Đã xem” hiện dưới tin cuối của bạn khi người khác đã đọc tới đó (nhóm: kèm số người đã xem).",
        "Số tin chưa đọc cập nhật tức thời trên từng hội thoại và trên biểu tượng “Tin nhắn” ở thanh trên.",
        "Nếu tạm mất kết nối, hệ thống tự kết nối lại.",
    ])
    m.figure("chat_32_dang_soan_tin", "Dòng “… đang soạn tin…”, tin hệ thống, biên nhận “Đã xem (2)” và dấu QTV")
    m.figure("chat_33_truc_tuyen", "Trạng thái trực tuyến của các đồng chí trong hội thoại")
    m.p("Về kỹ thuật: kênh dùng WebSocket nội bộ (đường dẫn /chats/ws), xác thực bằng phiên đăng nhập "
        "hiện tại, chạy chung một tiến trình với máy chủ — không cần dịch vụ ngoài.", italic=True, color=GRAYTXT)
    m.pagebreak()

    # =============================================================== 14
    m.h(1, "14. TIN HỆ THỐNG")
    m.p("Một số thay đổi về nhóm được ghi lại ngay trong dòng hội thoại dưới dạng “tin hệ thống” — "
        "chữ xám, căn giữa, không có bong bóng và không tính là tin chưa đọc:")
    m.bullets([
        "“<A> đã thêm <B> vào nhóm”",
        "“<A> đã đưa <B> ra khỏi nhóm” / “<A> đã rời nhóm”",
        "“<A> đổi tên nhóm thành «…»”",
    ])
    m.figure("chat_30_tin_he_thong", "Tin hệ thống ghi lại việc đổi tên nhóm ngay trong dòng hội thoại")
    m.pagebreak()

    # =============================================================== 15
    m.h(1, "15. QUẢN TRỊ NHÓM")
    m.p("Mở cột phải (nút chữ i). Tuỳ quyền, bạn sẽ thấy các công cụ quản trị nhóm.")

    m.h(2, "15.1. Đổi tên nhóm")
    m.steps([
        "Bấm biểu tượng bút chì cạnh tên nhóm ở đầu cột phải.",
        "Sửa tên trong ô, bấm “Lưu”. Hệ thống sinh một “tin hệ thống” báo cả nhóm biết ai đã đổi tên.",
    ])
    m.p("Quyền: Quản trị viên nhóm hoặc Ban chỉ huy.", italic=True, color=GRAYTXT)
    m.figure("chat_28_doi_ten_nhom", "Đổi tên nhóm ngay tại cột thông tin hội thoại")

    m.h(2, "15.2. Phong / gỡ Quản trị viên nhóm (QTV)")
    m.p("Trong danh sách “Thành viên”, cạnh mỗi người (trừ người tạo nhóm) có biểu tượng KHIÊN. Bấm để "
        "phong người đó làm QTV nhóm; bấm lại để gỡ. Người có dấu “QTV” được thêm thành viên, ghim tin, "
        "thu hồi tin sai của người khác trong nhóm đó.")
    m.p("Quyền phong/gỡ QTV: NGƯỜI TẠO NHÓM hoặc Ban chỉ huy. Không đổi được quyền của chính người tạo nhóm.",
        italic=True, color=GRAYTXT)
    m.figure("chat_29_quan_tri_thanh_vien", "Danh sách thành viên: dấu “QTV”, biểu tượng khiên để phong/gỡ, hai nút nhanh Tắt thông báo / Lưu trữ")

    m.h(2, "15.3. Xoá thành viên · Tự rời nhóm")
    m.steps([
        "Trong danh sách “Thành viên”, bấm dấu “X” cạnh đồng chí cần loại khỏi nhóm → xác nhận.",
        "Mọi thành viên đều có thể TỰ rời nhóm. Chỉ người tạo nhóm hoặc Quản trị hệ thống mới xoá được người khác.",
    ])
    m.figure("chat_16_drawer_quan_tri", "Cột thông tin nhóm: nút xoá từng thành viên (X) và nút “Xoá nhóm”")

    m.h(2, "15.4. Xoá nhóm")
    m.steps([
        "Bấm “Xoá nhóm (xoá cả tin nhắn & tệp)”.",
        "Đọc kỹ cảnh báo rồi xác nhận.",
    ])
    m.figure("chat_17_xac_nhan_xoa_nhom", "Hộp xác nhận trước khi xoá nhóm")
    m.note("Xoá nhóm là XOÁ CỨNG: toàn bộ tin nhắn và tệp đính kèm của nhóm bị xoá vĩnh viễn, không "
           "khôi phục. Quyền xoá nhóm: người tạo nhóm HOẶC Quản trị hệ thống. Thao tác được ghi Nhật ký an ninh.")
    m.pagebreak()

    # =============================================================== 16
    m.h(1, "16. TẮT THÔNG BÁO & LƯU TRỮ HỘI THOẠI")
    m.p("Hai tuỳ chọn này CHỈ áp dụng cho riêng bạn, không ảnh hưởng người khác.")
    m.h(2, "16.1. Tắt thông báo")
    m.p("Bấm nút CON MẮT GẠCH ở góc phải khung chat (hoặc nút “Tắt thông báo” ở cột phải). Hội thoại đó "
        "sẽ có biểu tượng 🔕 trong danh sách và không làm nổi số chưa đọc. Bấm lại để bật thông báo.")
    m.figure("chat_34_tat_thong_bao", "Hội thoại đã tắt thông báo mang biểu tượng 🔕 trong danh sách")
    m.h(2, "16.2. Lưu trữ")
    m.p("Bấm nút “Lưu trữ” ở cột phải để ẩn hội thoại khỏi danh sách chính (dùng cho các nhóm ít hoạt "
        "động). Xem lại bằng thẻ lọc “Lưu trữ” ở cột trái; mở ra và bấm “Bỏ lưu trữ” để đưa về danh sách chính.")
    m.figure("chat_35_luu_tru", "Thẻ lọc “Lưu trữ” liệt kê các hội thoại đã lưu trữ")
    m.pagebreak()

    # =============================================================== 17
    m.h(1, "17. PHÂN QUYỀN TÓM TẮT")
    m.table(
        ["Thao tác", "Ai được làm"],
        [
            ["Nhắn tin 1-1", "Mọi tài khoản đã kích hoạt"],
            ["Lập nhóm trực tiếp (dùng ngay)", "Ban chỉ huy: vai trò 0, 1, 2, 3"],
            ["Lập nhóm (phải chờ duyệt)", "Cán bộ / Người dùng: vai trò 4, 5"],
            ["Duyệt / từ chối nhóm chờ", "Quản trị hệ thống (0), Lữ trưởng – Chính uỷ (1)"],
            ["Gửi tin / gửi tệp / trả lời / cảm xúc / chuyển tiếp", "Thành viên của hội thoại (nhóm đã được duyệt)"],
            ["Sửa tin", "Chỉ người gửi tin đó"],
            ["Thu hồi tin", "Người gửi, hoặc QTV nhóm, hoặc Ban chỉ huy (0–3)"],
            ["Ghim / bỏ ghim tin", "Nhóm: QTV nhóm hoặc Ban chỉ huy. Chat 1-1: cả hai"],
            ["Đổi tên nhóm", "QTV nhóm hoặc Ban chỉ huy"],
            ["Thêm thành viên vào nhóm", "Người tạo nhóm, QTV nhóm hoặc Ban chỉ huy (0–3)"],
            ["Phong / gỡ QTV nhóm", "Người tạo nhóm hoặc Ban chỉ huy"],
            ["Tự rời nhóm", "Mọi thành viên"],
            ["Xoá thành viên khác", "Người tạo nhóm hoặc Quản trị hệ thống (0)"],
            ["Xoá nhóm", "Người tạo nhóm hoặc Quản trị hệ thống (0)"],
            ["Tắt thông báo / Lưu trữ", "Từng người, cho riêng mình"],
        ],
        widths=[3.3, 3.3],
    )
    m.pagebreak()

    # =============================================================== 18
    m.h(1, "18. AN TOÀN THÔNG TIN & NHẬT KÝ")
    m.bullets([
        "Toàn bộ hội thoại, tin nhắn, tệp đính kèm nằm trong cơ sở dữ liệu và ổ đĩa của máy chủ đơn vị — không đi ra ngoài.",
        "Khu vực Tin nhắn Tác chiến gắn nhãn “MẬT – MẠNG QUÂN SỰ LAN NỘI BỘ”; chỉ tài khoản do đơn vị cấp mới truy cập.",
        "Tệp tải lên bị quét chữ ký nhị phân, chặn tệp thực thi và mã độc / web shell trước khi lưu.",
        "Nhật ký an ninh ghi lại: tạo nhóm, xin tạo nhóm, duyệt nhóm, từ chối nhóm (kèm lý do), đổi tên nhóm, xoá nhóm — gồm người thực hiện, thời điểm, đối tượng.",
        "Tin hệ thống trong dòng hội thoại giúp mọi thành viên tự thấy thay đổi về thành viên / tên nhóm.",
        "Tài khoản đăng nhập sai nhiều lần bị khoá tạm thời (chống dò mật khẩu).",
    ])
    m.pagebreak()

    # =============================================================== 19
    m.h(1, "19. KẾT LUẬN & ĐỀ XUẤT")
    m.p("Kênh Tin nhắn Tác chiến nội bộ (bản nâng cấp) đưa việc trao đổi công việc của đơn vị về đúng "
        "một môi trường do đơn vị làm chủ: chạy trong mạng nội bộ, phân quyền rõ ràng, có luồng duyệt "
        "nhóm, có lưu vết; đầy đủ thao tác trên tin nhắn (trả lời, sửa, thu hồi, cảm xúc, ghim, chuyển "
        "tiếp), có tìm kiếm, tắt thông báo, lưu trữ và nhận tin thời gian thực (đang soạn tin, trực "
        "tuyến, đã xem) — với trải nghiệm quen thuộc như các ứng dụng nhắn tin phổ biến.")
    m.p("Đề xuất Chỉ huy đơn vị:", bold=True)
    m.bullets([
        "Phổ biến toàn đơn vị SỬ DỤNG kênh này thay cho nhắn tin qua mạng công cộng đối với nội dung công việc.",
        "Quy định thống nhất việc lập các nhóm kíp trực / nhóm công tác và giao Ban Tham mưu (hoặc bộ phận được "
        "phân công) làm đầu mối quản trị, duyệt nhóm, phong QTV nhóm.",
        "Giao bộ phận kỹ thuật bảo đảm máy chủ chạy liên tục, sao lưu định kỳ (đã có công cụ sẵn trong phần mềm).",
    ])
    m.p(" ")
    m.p("NGƯỜI BIÊN SOẠN", bold=True)
    m.p("Lê Văn Quỳnh – Đại đội 5, Lữ đoàn Thông tin 21, Bộ đội Biên phòng")

    m.save()
    print(f"Da xuat: {OUT}")


if __name__ == "__main__":
    build()
